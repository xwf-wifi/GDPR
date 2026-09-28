"""Build a review workbook from local evidence, preserving human edits on refresh."""

import argparse
import io
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image, ImageOps


SITE_SHEET = "网站总表"
COOKIE_SHEET = "Cookie明细"
SITE_HEADERS = [
    "排名", "域名", "采集状态", "HTTP", "Cookie数", "截图", "政策链接", "GDPR声明状态",
    "声明原文短摘录", "声明来源链接", "核查日期", "同意界面", "行为判断", "判断依据",
]
COOKIE_HEADERS = [
    "排名", "域名", "场景", "Cookie名称", "Cookie域名", "路径", "HttpOnly", "Secure",
    "SameSite", "用途", "用途证据链接", "分类状态",
]
SITE_EDITABLE = SITE_HEADERS[6:]
COOKIE_EDITABLE = COOKIE_HEADERS[9:]
HEADER_FILL = PatternFill("solid", fgColor="20344D")
INPUT_FILL = PatternFill("solid", fgColor="FFF3D8")


def saved_edits(path):
    if not path.exists():
        return {}, {}
    workbook = load_workbook(path, read_only=True, data_only=False)
    saved_sites, saved_cookies = {}, {}
    for name, key_fields, fields, target in (
        (SITE_SHEET, SITE_HEADERS[:2], SITE_EDITABLE, saved_sites),
        (COOKIE_SHEET, COOKIE_HEADERS[:6], COOKIE_EDITABLE, saved_cookies),
    ):
        if name not in workbook:
            continue
        rows = workbook[name].iter_rows(values_only=True)
        headers = next(rows, ())
        positions = {header: index for index, header in enumerate(headers)}
        if not all(field in positions for field in (*key_fields, *fields)):
            continue
        for values in rows:
            key = tuple(values[positions[field]] for field in key_fields)
            target[key] = {field: values[positions[field]] for field in fields}
    workbook.close()
    return saved_sites, saved_cookies


def style_sheet(sheet, headers, widths, editable_start):
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(2, sheet.max_row)}"
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 34
    for col, (header, width) in enumerate(zip(headers, widths), start=1):
        sheet.column_dimensions[get_column_letter(col)].width = width
        cell = sheet.cell(1, col)
        cell.fill = HEADER_FILL
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for col, cell in enumerate(row, start=1):
            cell.font = Font(name="Microsoft YaHei", size=10, color="202B38")
            cell.alignment = Alignment(vertical="center", wrap_text=col >= editable_start)
            if col >= editable_start:
                cell.fill = INPUT_FILL


def add_thumbnail(sheet, path, row):
    if not path.exists():
        return False
    with Image.open(path) as source:
        thumb = ImageOps.contain(source.convert("RGB"), (190, 110))
        stream = io.BytesIO()
        thumb.save(stream, format="JPEG", quality=78)
    stream.seek(0)
    image = ExcelImage(stream)
    image.anchor = f"F{row}"
    sheet.add_image(image)
    sheet.row_dimensions[row].height = 88
    return True


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--sites", type=Path, default=Path("data/top100.json"))
    parser.add_argument("--results", type=Path, default=Path("results"))
    parser.add_argument("--output", type=Path, default=Path("reports/GDPR_review.xlsx"))
    args = parser.parse_args()

    sites = json.loads(args.sites.read_text(encoding="utf-8"))["sites"]
    site_edits, cookie_edits = saved_edits(args.output)
    workbook = Workbook()
    overview = workbook.active
    overview.title = SITE_SHEET
    details = workbook.create_sheet(COOKIE_SHEET)
    overview.append(SITE_HEADERS)
    details.append(COOKIE_HEADERS)
    screenshots = collected = cookie_rows = 0

    for site in sites:
        rank, domain = site["rank"], site["domain"]
        folder = args.results / f"{rank:03d}_{domain}"
        evidence_path = folder / "evidence.json"
        evidence = json.loads(evidence_path.read_text(encoding="utf-8")) if evidence_path.exists() else None
        snapshot = (evidence or {}).get("snapshot") or {}
        cookie_list = snapshot.get("cookies") or []
        status = "未采集"
        if evidence:
            status = "采集失败" if evidence.get("error") or not snapshot else "首次访问已采集"
            collected += status == "首次访问已采集"
        edits = site_edits.get((rank, domain), {})
        row = [rank, domain, status, (evidence or {}).get("http_status"),
               len(cookie_list) if snapshot else None, ""] + [edits.get(field) for field in SITE_EDITABLE]
        if not edits.get("GDPR声明状态"):
            row[SITE_HEADERS.index("GDPR声明状态")] = "未核查"
        if not edits.get("行为判断"):
            row[SITE_HEADERS.index("行为判断")] = "证据不足"
        overview.append(row)
        if status == "首次访问已采集" and add_thumbnail(overview, folder / "initial.png", overview.max_row):
            overview.cell(overview.max_row, SITE_HEADERS.index("截图") + 1).value = "首页截图"
            screenshots += 1
        for cookie in cookie_list:
            key = (rank, domain, evidence.get("scenario"), cookie.get("name"),
                   cookie.get("domain"), cookie.get("path"))
            annotations = cookie_edits.get(key, {})
            details.append(list(key) + [cookie.get("httpOnly"), cookie.get("secure"),
                           cookie.get("sameSite"), annotations.get("用途"),
                           annotations.get("用途证据链接"), annotations.get("分类状态") or "未核查"])
            cookie_rows += 1

    style_sheet(overview, SITE_HEADERS, [8,22,18,9,11,29,34,18,44,34,15,20,18,42], 7)
    style_sheet(details, COOKIE_HEADERS, [8,22,19,28,27,15,13,13,16,36,34,17], 10)
    for sheet in (overview, details):
        sheet.auto_filter.ref = f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"
    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)
    print(f"已生成 {args.output}: {len(sites)} 个网站、{collected} 个已采集、{cookie_rows} 条 Cookie、{screenshots} 张截图")
    print("保留了旧工作簿中的人工填写栏；不导出 Cookie 值。请先核对截图与声明，再分享工作簿。")


if __name__ == "__main__":
    main()
