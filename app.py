"""Small Windows desktop interface for collecting and reviewing Tranco sites."""

import json
import os
import queue
import sys
import threading
import tkinter as tk
from pathlib import Path
from tkinter import messagebox, ttk
from urllib.parse import urlsplit

from openpyxl import load_workbook
from PIL import Image, ImageTk
from export_workbook import COOKIE_HEADERS, COOKIE_SHEET, SITE_HEADERS, SITE_SHEET, build_workbook


ROOT = Path(sys.executable).resolve().parent if getattr(sys, "frozen", False) else Path(__file__).resolve().parent
SITES = ROOT / "data" / "top100.json"
RESULTS = ROOT / "results"
REPORT = ROOT / "reports" / "GDPR_review.xlsx"


def resolve_site(value, sites):
    """Accept a domain or homepage URL; only the fixed Tranco sample is eligible."""
    raw = value.strip()
    if not raw:
        raise ValueError("请输入网站域名或网址。")
    parsed = urlsplit(raw if "://" in raw else "https://" + raw)
    if parsed.scheme not in {"http", "https"} or parsed.username or parsed.password:
        raise ValueError("请输入普通的 http(s) 网站网址。")
    host = (parsed.hostname or "").lower().rstrip(".")
    domain = host[4:] if host.startswith("www.") else host
    matched = next((item for item in sites if item["domain"] == domain), None)
    if not matched:
        raise ValueError("这个域名不在当前 Tranco 前 100 名中，请核对拼写。")
    return matched


def saved_row(sheet, headers, rank, domain):
    for row in sheet.iter_rows(min_row=2):
        if row[0].value == rank and row[1].value == domain:
            return {name: row[index].value for index, name in enumerate(headers)}
    return {}


class ResearchApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Cookie 核查工具")
        self.root.geometry("1080x760")
        self.root.minsize(900, 620)
        self.sites = json.loads(SITES.read_text(encoding="utf-8"))["sites"]
        self.events = queue.Queue()
        self.site = None
        self.preview = None
        self.cookie_key = None

        top = ttk.Frame(root, padding=14)
        top.pack(fill="x")
        ttk.Label(top, text="网站网址：").pack(side="left")
        self.url = ttk.Entry(top, width=45)
        self.url.pack(side="left", padx=6)
        self.url.insert(0, "baidu.com")
        self.collect_button = ttk.Button(top, text="采集网站", command=self.start_collect)
        self.collect_button.pack(side="left", padx=6)
        ttk.Button(top, text="查看已有结果", command=self.show_site).pack(side="left", padx=6)
        ttk.Button(top, text="打开工作簿", command=self.open_report).pack(side="left", padx=6)

        self.status = tk.StringVar(value="输入 Tranco 名单中的域名，然后点击“采集网站”。")
        ttk.Label(root, textvariable=self.status, padding=(14, 0)).pack(fill="x")
        body = ttk.Frame(root, padding=14)
        body.pack(fill="both", expand=True)
        left = ttk.Frame(body)
        left.pack(side="left", fill="both", expand=True)
        right = ttk.Frame(body)
        right.pack(side="right", fill="y", padx=(16, 0))

        self.title = tk.StringVar(value="尚未选择网站")
        ttk.Label(left, textvariable=self.title, font=("Microsoft YaHei", 14, "bold")).pack(anchor="w")
        form = ttk.LabelFrame(left, text="人工核查记录（填写后点击保存）", padding=10)
        form.pack(fill="both", expand=True, pady=10)
        self.fields = {}
        for label in ("政策链接", "GDPR声明状态", "声明来源链接", "核查日期", "同意界面", "行为判断"):
            line = ttk.Frame(form)
            line.pack(fill="x", pady=4)
            ttk.Label(line, text=label, width=15).pack(side="left")
            if label in ("GDPR声明状态", "行为判断"):
                values = ("未核查", "明确声称", "未发现明确声明", "不确定") if label == "GDPR声明状态" else ("证据不足", "待核查", "潜在问题", "未发现问题")
                widget = ttk.Combobox(line, values=values, state="readonly", width=42)
            else:
                widget = ttk.Entry(line, width=58)
            widget.pack(side="left", fill="x", expand=True)
            self.fields[label] = widget
        for label in ("声明原文短摘录", "判断依据"):
            ttk.Label(form, text=label).pack(anchor="w", pady=(7, 0))
            widget = tk.Text(form, height=3, wrap="word")
            widget.pack(fill="x")
            self.fields[label] = widget
        ttk.Button(form, text="保存人工记录", command=self.save_review).pack(anchor="e", pady=8)

        cookie_frame = ttk.LabelFrame(left, text="Cookie 明细（选中一条后填写）", padding=10)
        cookie_frame.pack(fill="both", expand=True)
        self.cookie_tree = ttk.Treeview(cookie_frame, columns=("name", "domain"), show="headings", height=5)
        self.cookie_tree.heading("name", text="名称")
        self.cookie_tree.heading("domain", text="域名")
        self.cookie_tree.column("name", width=180)
        self.cookie_tree.column("domain", width=220)
        self.cookie_tree.pack(fill="both", expand=True)
        self.cookie_tree.bind("<<TreeviewSelect>>", self.select_cookie)
        self.cookie_inputs = {}
        for label in ("用途", "用途证据链接", "分类状态"):
            line = ttk.Frame(cookie_frame)
            line.pack(fill="x", pady=2)
            ttk.Label(line, text=label, width=13).pack(side="left")
            widget = (ttk.Combobox(line, values=("未核查", "必要", "可选", "不确定"), state="readonly")
                      if label == "分类状态" else ttk.Entry(line))
            widget.pack(side="left", fill="x", expand=True)
            self.cookie_inputs[label] = widget
        self.cookie_purpose = self.cookie_inputs["用途"]
        self.cookie_source = self.cookie_inputs["用途证据链接"]
        self.cookie_class = self.cookie_inputs["分类状态"]
        ttk.Button(cookie_frame, text="保存所选 Cookie", command=self.save_cookie).pack(anchor="e", pady=5)

        ttk.Label(right, text="首次访问截图", font=("Microsoft YaHei", 11, "bold")).pack(anchor="w")
        self.image_label = ttk.Label(right, text="采集后在这里预览", width=42)
        self.image_label.pack(pady=10)
        ttk.Label(right, text="截图仅供核查页面状态；\n不能单凭截图判定合规。", justify="left").pack(anchor="w")
        self.root.after(150, self.poll_events)

    def chosen_site(self):
        try:
            return resolve_site(self.url.get(), self.sites)
        except ValueError as exc:
            messagebox.showerror("网址无效", str(exc))
            return None

    def start_collect(self):
        site = self.chosen_site()
        if not site:
            return
        self.collect_button.configure(state="disabled")
        self.status.set(f"正在采集 {site['domain']}，请稍候……")
        threading.Thread(target=self.collect_worker, args=(site,), daemon=True).start()

    def collect_worker(self, site):
        try:
            from playwright.sync_api import sync_playwright
            from collect import collect_site
            with sync_playwright() as playwright:
                browser = playwright.chromium.launch(headless=True)
                try:
                    result = collect_site(browser, site, RESULTS, 20000)
                finally:
                    browser.close()
            build_workbook(SITES, RESULTS, REPORT)
            self.events.put(("done", site, result))
        except Exception as exc:
            self.events.put(("error", site, str(exc)))

    def poll_events(self):
        try:
            while True:
                kind, site, payload = self.events.get_nowait()
                self.collect_button.configure(state="normal")
                if kind == "error":
                    self.status.set("采集失败，请检查依赖或网络。")
                    messagebox.showerror("采集失败", payload)
                else:
                    self.load_site(site)
                    self.status.set(f"{site['domain']}：{'采集失败，请看证据文件' if payload.get('error') else '采集完成，可填写核查记录'}")
        except queue.Empty:
            pass
        self.root.after(150, self.poll_events)

    def show_site(self):
        site = self.chosen_site()
        if site:
            if not REPORT.exists():
                build_workbook(SITES, RESULTS, REPORT)
            self.load_site(site)

    def load_site(self, site):
        self.site = site
        self.title.set(f"第 {site['rank']} 名：{site['domain']}")
        wb = load_workbook(REPORT, read_only=True, data_only=True)
        data = saved_row(wb[SITE_SHEET], SITE_HEADERS, site["rank"], site["domain"])
        for name, widget in self.fields.items():
            value = str(data.get(name) or "")
            if isinstance(widget, tk.Text):
                widget.delete("1.0", "end")
                widget.insert("1.0", value)
            elif isinstance(widget, ttk.Combobox):
                widget.set(value)
            else:
                widget.delete(0, "end")
                widget.insert(0, value)
        self.cookie_tree.delete(*self.cookie_tree.get_children())
        self.cookie_rows = {}
        for row in wb[COOKIE_SHEET].iter_rows(min_row=2, values_only=True):
            if row[0] == site["rank"] and row[1] == site["domain"]:
                item = self.cookie_tree.insert("", "end", values=(row[3], row[4]))
                self.cookie_rows[item] = row
        wb.close()
        image_path = RESULTS / f"{site['rank']:03d}_{site['domain']}" / "initial.png"
        if image_path.exists() and data.get("采集状态") == "首次访问已采集":
            with Image.open(image_path) as original:
                copy = original.copy()
            copy.thumbnail((340, 220))
            self.preview = ImageTk.PhotoImage(copy)
            self.image_label.configure(image=self.preview, text="")
        else:
            self.image_label.configure(image="", text="暂无有效截图")
            self.preview = None

    def select_cookie(self, _event):
        selected = self.cookie_tree.selection()
        if not selected:
            return
        row = self.cookie_rows[selected[0]]
        self.cookie_key = tuple(row[:6])
        for widget, value in ((self.cookie_purpose, row[9]), (self.cookie_source, row[10])):
            widget.delete(0, "end")
            widget.insert(0, value or "")
        self.cookie_class.set(row[11] or "未核查")

    def save_review(self):
        if not self.site:
            messagebox.showinfo("先选择网站", "请先输入网址并查看或采集网站。")
            return
        try:
            wb = load_workbook(REPORT)
            for row in wb[SITE_SHEET].iter_rows(min_row=2):
                if (row[0].value, row[1].value) == (self.site["rank"], self.site["domain"]):
                    for name, widget in self.fields.items():
                        value = widget.get("1.0", "end").strip() if isinstance(widget, tk.Text) else widget.get().strip()
                        row[SITE_HEADERS.index(name)].value = value
                    break
            wb.save(REPORT)
            self.status.set("人工记录已保存到工作簿。")
        except PermissionError:
            messagebox.showerror("无法保存", "请先关闭 Excel 中打开的工作簿，再点击保存。")

    def save_cookie(self):
        if not self.cookie_key:
            messagebox.showinfo("先选 Cookie", "请先选中一条 Cookie。")
            return
        try:
            wb = load_workbook(REPORT)
            for row in wb[COOKIE_SHEET].iter_rows(min_row=2):
                if tuple(cell.value for cell in row[:6]) == self.cookie_key:
                    row[COOKIE_HEADERS.index("用途")].value = self.cookie_purpose.get().strip()
                    row[COOKIE_HEADERS.index("用途证据链接")].value = self.cookie_source.get().strip()
                    row[COOKIE_HEADERS.index("分类状态")].value = self.cookie_class.get()
                    break
            wb.save(REPORT)
            self.status.set("Cookie 用途记录已保存。")
        except PermissionError:
            messagebox.showerror("无法保存", "请先关闭 Excel 中打开的工作簿，再点击保存。")

    def open_report(self):
        if not REPORT.exists():
            build_workbook(SITES, RESULTS, REPORT)
        if sys.platform == "win32":
            os.startfile(REPORT)
        else:
            messagebox.showinfo("工作簿位置", str(REPORT))


def main():
    root = tk.Tk()
    ResearchApp(root)
    if "--smoke" in sys.argv:
        root.update()
        root.destroy()
        return
    root.mainloop()


if __name__ == "__main__":
    main()
